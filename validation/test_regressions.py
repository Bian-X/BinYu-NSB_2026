"""Regression tests using synthetic data and temporary metadata files."""
import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import numpy as np
import pandas as pd
from openpyxl import load_workbook

REPO=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(REPO/'QC_analysis/scripts'))
import qc_analysis as qa
from metadata_io import read_sample_metadata,match_metadata
from figure_2F_injection_order import build_figure_2f_correlation
spec=importlib.util.spec_from_file_location('pca_fixed',REPO/'sample_analysis/scripts/python/brain_region_pca.py')
pca=importlib.util.module_from_spec(spec);spec.loader.exec_module(pca)


class MetadataTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder=Path(self.temp.name)
        self.table=pd.DataFrame({'raw_file':['Astral_MSP2600441_1_G1_1_PL.raw','Astral_MSP2600441_QC_1.raw'],
             'sample_type':['BIO','QC'],'region':['PL','QC'],'group':['G1','QC'],'injection_order':[1,2]})

    def write(self,table,name='metadata.xlsx',startrow=0):
        path=self.folder/name
        if path.suffix=='.csv':table.to_csv(path,index=False)
        else:table.to_excel(path,index=False,startrow=startrow)
        return path

    def test_single_grouped_and_csv_headers_agree(self):
        a=read_sample_metadata(self.write(self.table))
        path=self.write(self.table,'grouped.xlsx',1)
        wb=load_workbook(path);wb.active.cell(1,1,'sample-data');wb.save(path)
        b=read_sample_metadata(path)
        c=read_sample_metadata(self.write(self.table,'metadata.csv'))
        pd.testing.assert_frame_equal(a,b)
        pd.testing.assert_frame_equal(a,c)

    def test_duplicate_order_rejected(self):
        t=self.table.copy();t.injection_order=[1,1]
        with self.assertRaisesRegex(ValueError,'duplicate'):read_sample_metadata(self.write(t))

    def test_invalid_order_rejected(self):
        for value in [0,-1,1.5,np.nan,np.inf]:
            with self.subTest(value=value):
                t=self.table.copy();t.injection_order=[value,2]
                with self.assertRaises(ValueError):read_sample_metadata(self.write(t))

    def test_blank_biological_fields_rejected(self):
        for field in ['region','group']:
            t=self.table.copy();t.loc[0,field]='  '
            with self.assertRaisesRegex(ValueError,'lacks'):read_sample_metadata(self.write(t))

    def test_unknown_or_inconsistent_identity_rejected(self):
        for field,value in [('region','UNKNOWN'),('group','G99'),('group','G6'),('sample_type','QC')]:
            with self.subTest(field=field,value=value):
                t=self.table.copy();t.loc[0,field]=value
                with self.assertRaises(ValueError):read_sample_metadata(self.write(t))

    def test_sample_set_mismatch_rejected(self):
        t=read_sample_metadata(self.write(self.table))
        with self.assertRaisesRegex(ValueError,'mismatch'):match_metadata(t,t.sample_id.iloc[:1])

    def test_blank_qc_region_is_normalized_without_changing_biological_fields(self):
        t=self.table.copy();t.loc[1,'region']=np.nan
        actual=read_sample_metadata(self.write(t,'blank_qc.csv'))
        self.assertEqual(actual.loc[1,'region'],'QC')
        self.assertEqual(actual.loc[0,'region'],'PL')


class PreprocessingTests(unittest.TestCase):
    def test_qc_loader_uses_only_manifest_reports(self):
        with tempfile.TemporaryDirectory() as directory:
            folder=Path(directory)
            expected=pd.read_csv(REPO/'metadata/input_manifest.csv').filename.tolist()
            for filename in expected:
                (folder/filename).touch()
            (folder/'metadata.csv').write_text('raw_file,sample_type\nfixture,BIO\n',encoding='utf-8')
            with patch.object(qa,'load_report',side_effect=lambda path:path):
                files=qa.load_all_reports(folder)
            self.assertEqual([path.name for path in files],expected)
            alias=folder/f'{Path(expected[0]).stem} .csv'
            (folder/expected[0]).rename(alias)
            with patch.object(qa,'load_report',side_effect=lambda path:path):
                self.assertEqual(qa.load_all_reports(folder)[0],alias)
            alias.unlink()
            with patch.object(qa,'load_report',side_effect=lambda path:path):
                with self.assertRaisesRegex(FileNotFoundError,'Expected one report'):
                    qa.load_all_reports(folder)

    def test_g_group_header_is_used_without_relabeling(self):
        raw='[7] Astral_MSP2600441_253_G6_1_PVH.raw.PG.Quantity'
        sample='Astral_MSP2600441_253_G6_1_PVH.raw'
        self.assertEqual(qa.clean_quantity_column(raw),sample)
        self.assertEqual(qa.parse_sample(sample,'PVH.csv','PVH')['group'],'G6')
        self.assertEqual(qa.parse_sample('Astral_MSP2600441_7_G2_1_PL.raw','PL.csv','PL')['level'],'unpaired_tone_shock_control')
        invalid_group=sample.replace('_G6_','_E3_')
        self.assertEqual(qa.parse_sample(invalid_group,'PVH.csv','PVH')['sample_type'],'UNKNOWN')

    def test_qc_cv_and_regression_have_explicitly_different_populations(self):
        ann=pd.DataFrame({'PG.Organisms':['Mus musculus','Unknown',np.nan,'Mus musculus'],
                          'PG.ProteinDescriptions':['mouse','Tax_Id=10090','iRT','mouse']})
        values=pd.DataFrame([[1.,2.,3.],[1.,2.,3.],[1.,2.,3.],[1.,np.nan,np.nan]],index=['m','unknown','iRT-Kit_WR_fusion','sparse'])
        report=qa.ReportData(Path('QC.csv'),'QC',True,ann,values,pd.DataFrame())
        audit=qa.qc_regression_selection(report,min_detected=2)
        self.assertEqual(int(audit.CVEligible.sum()),3)
        self.assertEqual(int(audit.RegressionIncluded.sum()),1)
        self.assertEqual(audit.loc[audit.RegressionIncluded,'PG.ProteinAccessions'].tolist(),['m'])

    def test_auxiliary_pca_rejects_uncompleted_matrix(self):
        with self.assertRaisesRegex(ValueError,'median'):
            qa.compute_pca(pd.DataFrame([[1.,np.nan],[2.,3.]]))

    def test_exclusion_before_filter_and_exact_irt_removal(self):
        samples=[f'Astral_MSP2600441_{i}_G{g}_{r}_PVH' for i,(g,r) in enumerate(((g,r) for g in range(1,7) for r in range(1,5)),1)]
        samples[20]=pca.OUTLIER_SAMPLE
        records=[]
        for accession,observed in [('depends_on_excluded',samples[:19]+[samples[20]]),
                                   ('retained',samples[:20]),('Sirt1',samples[:20]),
                                   ('anchor',samples),('iRT-Kit_WR_fusion',samples)]:
            row=dict(zip(pca.ANNOTATION_COLUMNS,[accession,accession,'Mus musculus','fixture',2]))
            row.update({f'[{i+1}] {s}.raw.PG.Quantity':float(i+1) if s in observed else np.nan for i,s in enumerate(samples)})
            records.append(row)
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'PVH_protein_groups.csv';pd.DataFrame(records).to_csv(path,index=False)
            all_matrix,_,_=pca.read_and_preprocess_region(path)
            primary,_,qc=pca.read_and_preprocess_region(path,[pca.OUTLIER_SAMPLE])
            self.assertIn('depends_on_excluded',all_matrix.index)
            self.assertNotIn('depends_on_excluded',primary.index)
            self.assertNotIn(pca.OUTLIER_SAMPLE,primary.columns)
            self.assertEqual(qc['samples_retained'],23)
            self.assertIn('Sirt1',primary.index)
            self.assertNotIn('iRT-Kit_WR_fusion',primary.index)
            self.assertTrue(np.isfinite(primary.to_numpy()).all())
            invalid=pd.DataFrame(records).rename(columns={f'[1] {samples[0]}.raw.PG.Quantity':f'[1] {samples[0].replace("G1_1_PVH","G9_1_PL")}.raw.PG.Quantity'})
            invalid.to_csv(path,index=False)
            with self.assertRaisesRegex(ValueError,'identity|region'):pca.read_and_preprocess_region(path)

    def test_qc_views_share_one_matrix_and_drop_sparse_rows_first(self):
        m=pd.DataFrame([[1.,2,3,4,5],[2,4,6,8,np.nan],[3,5,7,11,13],[1000,np.nan,np.nan,np.nan,np.nan]],index=['a','b','c','sparse'])
        expected=qa.prepare_complete_matrix(m)
        self.assertNotIn('sparse',expected.index)
        self.assertTrue(np.isfinite(expected.to_numpy()).all())
        # A sparse row must not change normalization of retained features.
        pd.testing.assert_frame_equal(expected,qa.prepare_complete_matrix(m.drop(index='sparse')))
        with patch.object(qa,'load_all_reports',return_value=[]),patch.object(qa,'build_combined_matrix',return_value=(m,pd.DataFrame(),pd.DataFrame())):
            actual,_=build_figure_2f_correlation(Path('.'))
        pd.testing.assert_frame_equal(actual,expected.corr())

    def test_qc_metrics_never_fill_missing_values(self):
        table=pd.DataFrame({'Astral_MSP2600441_QC_1':[1.,np.nan],'Astral_MSP2600441_QC_2':[2.,2.]},index=['a','b'])
        meta=pd.DataFrame({'sample_id':table.columns,'sample_type':['QC','QC']})
        ann=pd.DataFrame({'PG.ProteinAccessions':['a','b']})
        cv=qa.qc_cv_table(table,meta,ann.set_index('PG.ProteinAccessions',drop=False)).set_index('PG.ProteinAccessions')
        self.assertEqual(cv.loc['b','qc_detected_n'],1)
        self.assertTrue(pd.isna(cv.loc['b','qc_cv_percent']))


if __name__=='__main__':unittest.main(verbosity=2)
